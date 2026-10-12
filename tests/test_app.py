import io, zipfile
import pytest
from docx import Document
from app import app, duration

@pytest.fixture
def client(tmp_path):
    app.config.update(TESTING=True,DATABASE=str(tmp_path/'test.sqlite'))
    with app.test_client() as c:
        c.get('/')
        with c.session_transaction() as s: token=s['csrf']
        c.environ_base['HTTP_X_CSRF_TOKEN']=token
        yield c

def officer(c,number='101',station='Central'):
    data=dict(regimental_number=number,rank='PC',name='Test Officer '+number,station=station,division='North',nis_number='TEST-NIS',bir_number='TEST-BIR',bank='Test Bank',telephone='123',cellular='456',account_number='TEST-ACCOUNT',address='Test Address')
    r=c.post('/api/officers',json=data); assert r.status_code==200
    return r.json['id']

def duty(ids,**extra):
    return dict(officer_ids=ids,date='2026-10-12',time_in='22:00',time_out='06:15',duty_type='Patrol',diary_in='1/2',diary_out='3/4',remarks='Test remark',**extra)

def test_minutes():
    assert duration('08:33','16:35')==482
    assert duration('22:00','06:15')==495
    with pytest.raises(ValueError):duration('08:00','08:00')

def test_batch_station_and_overlap_atomicity(client):
    a=officer(client);b=officer(client,'102','South')
    assert client.post('/api/duties',json=duty([a,b])).json['saved']==2
    rows=client.get('/api/duties?month=2026-10').json
    assert {r['station'] for r in rows}=={'Central','South'}
    assert all(r['minutes']==495 for r in rows)
    c=officer(client,'103')
    assert client.post('/api/duties',json=duty([c,a])).status_code==400
    assert len(client.get('/api/duties?month=2026-10').json)==2
    data=duty([a]);data.update(date='2026-10-13',time_in='05:00',time_out='07:00')
    assert client.post('/api/duties',json=data).status_code==400
    data.update(time_in='06:15',time_out='07:00',station='Override')
    assert client.post('/api/duties',json=data).status_code==200
    assert client.get('/api/duties?month=2026-10').json[0]['station']=='Override'

def test_export_template_and_zip(client):
    a=officer(client); b=officer(client,'102')
    client.post('/api/duties',json=duty([a]))
    r=client.get(f'/export/docx?month=2026-10&ids={a}')
    assert r.status_code==200
    doc=Document(io.BytesIO(r.data))
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    for expected in ['Test Officer 101','TEST-NIS','TEST-BIR','TEST-ACCOUNT','Test Bank','Test Address','8:15','Test remark','31.10.26']:
        assert expected in text
    for stale in ['JOHANNES','031485781','100006609-4','MARCH','ACP Criminal','9001']:
        assert stale not in text
    assert len(doc.tables[0].rows)==32
    pdf=client.get(f'/export/pdf?month=2026-10&ids={a}')
    assert pdf.status_code==200 and pdf.data.startswith(b'%PDF-')
    z=client.get(f'/export/docx?month=2026-10&ids={a},{b}')
    with zipfile.ZipFile(io.BytesIO(z.data)) as archive: assert len(archive.namelist())==2

def test_validation_and_persistence(client):
    a=officer(client)
    assert client.post('/api/officers',json=dict(regimental_number='101',rank='PC',name='Duplicate',station='X',division='X')).status_code==409
    assert client.post('/api/duties',json=duty([])).status_code==400
    bad=duty([a]);bad['date']='wrong'
    assert client.post('/api/duties',json=bad).status_code==400
    assert client.get('/export/pdf?ids=1&month=bad').status_code==400
    client.post('/api/duties',json=duty([a]))
    rows=client.get('/api/duties?month=2026-10').json
    assert len(rows)==1
    assert client.delete('/api/duties/'+str(rows[0]['id'])).status_code==200
    assert client.get('/api/duties?month=2026-10').json==[]

def test_csrf(client):
    assert client.post('/api/officers',json={},headers={'X-CSRF-Token':'bad'}).status_code==403


def test_bundled_template_has_no_sample_personal_data():
    from app import ROOT
    with zipfile.ZipFile(ROOT / 'assets/duty-template.docx') as archive:
        xml = '\n'.join(archive.read(name).decode() for name in archive.namelist() if name.endswith('.xml'))
    from lxml import etree
    with zipfile.ZipFile(ROOT / 'assets/duty-template.docx') as archive:
        visible = ''.join(etree.fromstring(archive.read('word/document.xml')).itertext())
    for stale in ['Johannes', 'JOHANNES', '031485781', '100006609-4', '9001', '653-8110', '480-6544', 'ROYAL BANK', 'ACP Criminal']:
        assert stale not in xml
        assert stale not in visible
    doc = Document(ROOT / 'assets/duty-template.docx')
    assert not doc.core_properties.author
    assert not doc.core_properties.last_modified_by
    assert all(not cell.text for row in doc.tables[0].rows[1:] for cell in row.cells)
