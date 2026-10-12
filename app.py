from pathlib import Path
from datetime import datetime, timedelta
import calendar, io, os, sqlite3, zipfile, secrets
from flask import Flask, request, jsonify, render_template, send_file, session, abort
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import legal
from xml.sax.saxutils import escape

ROOT = Path(__file__).parent
app = Flask(__name__)
app.config.update(SECRET_KEY=secrets.token_hex(32), DATABASE=os.environ.get('DUTY_DATABASE', str(ROOT/'instance/duties.sqlite')), MAX_CONTENT_LENGTH=1024*1024, SESSION_COOKIE_SAMESITE='Strict')
FIELDS = ['regimental_number','rank','name','nis_number','station','division','bir_number','bank','telephone','cellular','account_number','address']

def db():
    path = Path(app.config['DATABASE']); path.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(path); c.row_factory = sqlite3.Row; c.execute('PRAGMA foreign_keys=ON')
    c.execute('CREATE TABLE IF NOT EXISTS officers (id INTEGER PRIMARY KEY, '+','.join(f'{f} TEXT NOT NULL' for f in FIELDS)+', UNIQUE(regimental_number))')
    c.execute('CREATE TABLE IF NOT EXISTS duties (id INTEGER PRIMARY KEY, officer_id INTEGER NOT NULL REFERENCES officers(id), date TEXT NOT NULL, time_in TEXT NOT NULL, time_out TEXT NOT NULL, station TEXT NOT NULL, duty_type TEXT NOT NULL, minutes INTEGER NOT NULL, diary_in TEXT NOT NULL, diary_out TEXT NOT NULL, remarks TEXT NOT NULL, UNIQUE(officer_id,date,time_in,time_out))')
    return c

def duration(start, end):
    a=datetime.strptime(start,'%H:%M'); b=datetime.strptime(end,'%H:%M')
    if a==b: raise ValueError('Time in and time out must differ. Split shifts of 24 hours or more into separate entries.')
    if b<a: b+=timedelta(days=1)
    return int((b-a).total_seconds()/60)

def hours(n): return f'{n//60}:{n%60:02d}'

def clean(data, key, required=False):
    value=data.get(key,'')
    if not isinstance(value,str) or len(value)>500: raise ValueError(f'Invalid {key}. Maximum 500 characters.')
    value=value.strip()
    if required and not value: raise ValueError(f'{key.replace("_"," ").title()} is required.')
    return value

@app.before_request
def csrf():
    if request.method in ('POST','PUT','DELETE') and (not session.get('csrf') or request.headers.get('X-CSRF-Token')!=session['csrf']): abort(403)

@app.errorhandler(ValueError)
def bad(e): return jsonify(error=str(e)),400

@app.errorhandler(sqlite3.IntegrityError)
def conflict(e): return jsonify(error='This regimental number or duty entry already exists. No records were changed.'),409

@app.get('/')
def index():
    session.setdefault('csrf', secrets.token_hex(32))
    return render_template('index.html', token=session['csrf'])

@app.get('/api/officers')
def officers():
    with db() as c: return jsonify([dict(r) for r in c.execute('SELECT * FROM officers ORDER BY name')])

@app.route('/api/officers', methods=['POST'])
@app.route('/api/officers/<int:oid>', methods=['PUT'])
def save_officer(oid=None):
    data=request.get_json(); values=[clean(data,f,f in ('regimental_number','rank','name','station','division')) for f in FIELDS]
    with db() as c:
        if oid:
            if not c.execute('SELECT id FROM officers WHERE id=?',(oid,)).fetchone(): abort(404)
            c.execute('UPDATE officers SET '+','.join(f'{f}=?' for f in FIELDS)+' WHERE id=?', values+[oid])
        else: oid=c.execute('INSERT INTO officers ('+','.join(FIELDS)+') VALUES ('+','.join('?' for _ in FIELDS)+')',values).lastrowid
    return jsonify(id=oid)

def month_arg():
    month=request.args.get('month',datetime.now().strftime('%Y-%m'))
    try: datetime.strptime(month,'%Y-%m')
    except ValueError: raise ValueError('Choose a valid month.')
    if len(month)!=7: raise ValueError('Choose a valid month.')
    return month

@app.get('/api/duties')
def duties():
    month=month_arg()
    with db() as c: return jsonify([dict(r) for r in c.execute('SELECT d.*,o.name,o.rank,o.regimental_number FROM duties d JOIN officers o ON o.id=d.officer_id WHERE substr(date,1,7)=? ORDER BY date DESC,time_in DESC',(month,))])

@app.post('/api/duties')
def add_duty():
    data=request.get_json(); ids=data.get('officer_ids',[])
    if not isinstance(ids,list) or not ids or any(type(i)!=int for i in ids): raise ValueError('Select at least one officer.')
    ids=list(set(ids)); date=clean(data,'date',True)
    try: datetime.strptime(date,'%Y-%m-%d')
    except ValueError: raise ValueError('Choose a valid duty date.')
    tin=clean(data,'time_in',True); tout=clean(data,'time_out',True); minutes=duration(tin,tout)
    station=clean(data,'station'); kind=clean(data,'duty_type',True)
    with db() as c:
        for oid in ids:
            officer=c.execute('SELECT * FROM officers WHERE id=?',(oid,)).fetchone()
            if not officer: raise ValueError('An officer no longer exists.')
            start=datetime.fromisoformat(date+'T'+tin); end=start+timedelta(minutes=minutes)
            for old in c.execute('SELECT * FROM duties WHERE officer_id=?',(oid,)):
                a=datetime.fromisoformat(old['date']+'T'+old['time_in']); b=a+timedelta(minutes=old['minutes'])
                if start<b and end>a: raise ValueError(f'Duty overlaps an existing entry for {officer["name"]}. No entries were saved.')
            c.execute('INSERT INTO duties (officer_id,date,time_in,time_out,station,duty_type,minutes,diary_in,diary_out,remarks) VALUES (?,?,?,?,?,?,?,?,?,?)',(oid,date,tin,tout,station or officer['station'],kind,minutes,clean(data,'diary_in'),clean(data,'diary_out'),clean(data,'remarks')))
    return jsonify(saved=len(ids))

@app.delete('/api/duties/<int:did>')
def delete_duty(did):
    with db() as c: c.execute('DELETE FROM duties WHERE id=?',(did,))
    return jsonify(ok=True)

def report_rows(entries,month):
    year,m=map(int,month.split('-')); rows=[]
    for day in range(1,calendar.monthrange(year,m)[1]+1):
        date=f'{month}-{day:02}'
        found=[e for e in entries if e['date']==date]
        for e in found or [None]:
            rows.append([f'{day:02}.{m:02}.{year%100:02}',e['time_in'] if e else '',e['time_out'] if e else '',e['station'] if e else '',e['duty_type'] if e else '','',hours(e['minutes']) if e else '',e['diary_in'] if e else '',e['diary_out'] if e else '',e['remarks'] if e else ''])
    return rows

def profile_lines(o,month):
    return [f'MONTH: {datetime.strptime(month,"%Y-%m"):%B, %Y}    NO: {o["regimental_number"]}    RANK: {o["rank"]}    NAME: {o["name"]}',f'N.I.S NUMBER: {o["nis_number"]}    B.I.R NUMBER: {o["bir_number"]}',f'STATION/BRANCH: {o["station"]}    ADDRESS: {o["address"]}',f'DIVISION: {o["division"]}    COMMERCIAL BANK: {o["bank"]}',f'TEL. NO: {o["telephone"]}    CELL: {o["cellular"]}    ACCOUNT NO: {o["account_number"]}']

def word_report(o,entries,month):
    doc=Document(ROOT/'assets/duty-template.docx'); lines=profile_lines(o,month)
    for p in doc.paragraphs:
        t=p.text.strip()
        if t == 'DIVISION NAME': p.text=o['division'].upper()+' DIVISION'
        for prefix,line in zip(['MONTH:','N.I.S NUMBER:','STATION/BRANCH:','DIVISION:','TEL. NO:'],lines):
            if t.startswith(prefix): p.text=line
    table=doc.tables[0]
    for row in list(table.rows)[1:]: table._tbl.remove(row._tr)
    for values in report_rows(entries,month):
        cells=table.add_row().cells
        for cell,value in zip(cells,values):
            cell.text=value
            for p in cell.paragraphs:
                for run in p.runs: run.font.size=Pt(8)
    doc.add_paragraph('TOTAL HOURS (HH:MM): '+hours(sum(e['minutes'] for e in entries)))
    out=io.BytesIO(); doc.save(out); return out.getvalue()

def pdf_report(o,entries,month):
    out=io.BytesIO(); styles=getSampleStyleSheet(); small=styles['Normal']; small.fontSize=7; small.leading=9
    def p(s): return Paragraph(escape(str(s)),small)
    flow=[Paragraph('TRINIDAD &amp; TOBAGO POLICE SERVICE',styles['Heading2']),p('Office of the Adjutant Special Reserve Police'),p('7th Floor Riverside Plaza, Corner Besson and Picadilly Street, Port of Spain'),p('Telephone: 1(868) 612-2466'),Spacer(1,8),Paragraph('FULL TIME MONTHLY DUTY RETURN',styles['Heading3'])]
    flow += [p(s) for s in profile_lines(o,month)]; flow.append(Spacer(1,10))
    head=['Date','Time in','Time out','Station / Branch','Type of duty','Signature of sentry','Hours (HH:MM)','Diary in','Diary out','Remarks']
    table=Table([[p(v) for v in head]]+[[p(v) for v in r] for r in report_rows(entries,month)],colWidths=[36,32,32,68,58,50,42,42,42,54],repeatRows=1)
    table.setStyle(TableStyle([('LEFTPADDING',(0,0),(-1,-1),2),('RIGHTPADDING',(0,0),(-1,-1),2),('GRID',(0,0),(-1,-1),.35,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eef2')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
    certification=Table([[p('I certify that the hours stated were performed by me and are correct.'),p('Certified correct and corresponding with the Station Diary.'),p('Certified correct')],[p('_______________________'),p('_______________________'),p('_______________________')],[p('SIGNATURE'),p('INSPECTOR / ASST. SUPERINTENDENT'),p('SENIOR SUPERINTENDENT')]],colWidths=[185,185,186])
    certification.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,1),(-1,1),20)]))
    flow += [table,Spacer(1,10),p('TOTAL HOURS (HH:MM): '+hours(sum(e['minutes'] for e in entries))),Spacer(1,16),certification,Spacer(1,12),p('NOTE: This return is for one (1) month and must be submitted by the fifth (5th) of the following month.')]
    SimpleDocTemplate(out,pagesize=legal,leftMargin=28,rightMargin=28,topMargin=24,bottomMargin=24).build(flow)
    return out.getvalue()

@app.get('/export/<fmt>')
def export(fmt):
    if fmt not in ('pdf','docx'): abort(404)
    month=month_arg()
    try: ids=list(dict.fromkeys(int(i) for i in request.args.get('ids','').split(',') if i))
    except ValueError: raise ValueError('Invalid officer selection.')
    if not ids: raise ValueError('Select at least one officer.')
    outputs=[]
    with db() as c:
        for oid in ids:
            o=c.execute('SELECT * FROM officers WHERE id=?',(oid,)).fetchone()
            if not o: abort(404)
            entries=c.execute('SELECT * FROM duties WHERE officer_id=? AND substr(date,1,7)=? ORDER BY date,time_in',(oid,month)).fetchall()
            outputs.append((f'duty-{oid}-{month}.{fmt}',(pdf_report if fmt=='pdf' else word_report)(o,entries,month)))
    if len(outputs)==1:
        name,data=outputs[0]; return send_file(io.BytesIO(data),download_name=name,as_attachment=request.args.get('inline')!='1',mimetype='application/pdf' if fmt=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in outputs: z.writestr(name,data)
    out.seek(0); return send_file(out,download_name=f'duty-returns-{month}.zip',as_attachment=True)

if __name__=='__main__':
    app.run(host='127.0.0.1',port=int(os.environ.get('PORT','5000')),debug=False)
